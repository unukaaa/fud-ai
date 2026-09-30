#!/usr/bin/env python3
"""Refresh AUSNUT food measures without changing bundled food or nutrient data.

Requires openpyxl. The official FSANZ Food measures workbook is supplied with
--measures; it is not stored in this repository. The existing bundled JSON is
the nutrition input and output. No density row becomes a selectable portion.
"""

from __future__ import annotations

import argparse
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


HEADERS = (
    "Survey ID", "Public food key", "Food name", "Measure ID", "Quantity",
    "Descriptor 1\n", "Descriptor 2", "Descriptor 3", "Descriptor 4\n",
    "Gram amount", "Volume",
)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--measures", required=True, type=Path)
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    database = json.loads(args.database.read_text(encoding="utf-8"))
    foods = database["foods"]
    foods_by_id = {food["id"]: food for food in foods}
    if len(foods_by_id) != len(foods):
        raise ValueError("Bundled food IDs are not unique")

    workbook = load_workbook(args.measures, read_only=True, data_only=True)
    sheet = workbook["AUSNUT 2023"]
    if tuple(cell.value for cell in sheet[3][:11]) != HEADERS:
        raise ValueError("Unexpected official Food measures columns")

    measures_by_food: dict[str, list[dict]] = defaultdict(list)
    measure_ids: set[int] = set()
    total_rows = density_rows = 0
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        total_rows += 1
        survey_id = str(row[0])
        if survey_id not in foods_by_id:
            raise ValueError(f"Measure belongs to unknown Survey ID {survey_id}")
        descriptors = [row[index] for index in range(5, 9)]
        if str(descriptors[0] or "").strip().lower() == "density":
            density_rows += 1
            continue
        measure_id, quantity, grams, volume = row[3], row[4], row[9], row[10]
        if not isinstance(measure_id, int) or measure_id in measure_ids:
            raise ValueError(f"Missing or repeated Measure ID {measure_id}")
        if not all(isinstance(value, (int, float)) and math.isfinite(value) and value > 0
                   for value in (quantity, grams)):
            raise ValueError(f"Invalid quantity or gram amount at Measure ID {measure_id}")
        if not any(isinstance(part, str) and part.strip() for part in descriptors):
            raise ValueError(f"Missing descriptor at Measure ID {measure_id}")
        if not isinstance(volume, (int, float)) or not math.isfinite(volume):
            raise ValueError(f"Invalid source volume at Measure ID {measure_id}")
        measure_ids.add(measure_id)
        display_name = " ".join(" ".join(part.split()) for part in descriptors if part)
        measures_by_food[survey_id].append({
            "n": display_name, "q": quantity, "g": grams,
            "mid": measure_id, "d": descriptors, "v": volume,
        })

    if total_rows != density_rows + len(measure_ids):
        raise ValueError("Official measure rows were not completely classified")
    for food in foods:
        measures = measures_by_food.get(food["id"], [])
        if measures:
            food["measures"] = measures
        else:
            food.pop("measures", None)

    args.output.write_text(
        json.dumps(database, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    print(f"foods={len(foods)} rows={total_rows} density={density_rows} "
          f"eligible={len(measure_ids)} output={args.output}")


if __name__ == "__main__":
    main()
